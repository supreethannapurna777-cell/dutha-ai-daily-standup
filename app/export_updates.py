from pathlib import Path
import csv
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


DATABASE_PATH = Path("data/standup.db")
CSV_REPORT_PATH = Path("data/daily_updates.csv")
EXCEL_REPORT_PATH = Path("data/daily_updates.xlsx")


HEADERS = [
    "Date",
    "Response Time",
    "Team Member",
    "Team/Department",
    "Today's Tasks",
    "People to Connect With",
    "Blockers",
    "Dependencies",
    "Expected Completion",
    "Original Reply",
    "Reminder Status",
]


def get_update_rows():
    """Read processed updates from SQLite."""

    with sqlite3.connect(DATABASE_PATH) as connection:
        return connection.execute(
            """
            SELECT
                incoming_messages.received_at,
                processed_updates.sender_name,
                processed_updates.tasks,
                processed_updates.people_to_connect,
                processed_updates.blockers,
                processed_updates.dependencies,
                processed_updates.expected_completion,
                incoming_messages.original_reply
            FROM processed_updates
            JOIN incoming_messages
                ON processed_updates.message_id = incoming_messages.id
            ORDER BY incoming_messages.received_at
            """
        ).fetchall()


def build_report_rows(rows):
    """Convert database rows into report rows."""

    report_rows = []

    for row in rows:
        received_at = row[0]
        date, response_time = received_at.split("T", maxsplit=1)

        report_rows.append(
            [
                date,
                response_time,
                row[1],
                "Not specified",
                row[2],
                row[3],
                row[4],
                row[5],
                row[6],
                row[7],
                "Not started",
            ]
        )

    return report_rows


def export_csv(report_rows):
    """Export the updates to CSV."""

    with CSV_REPORT_PATH.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(HEADERS)
        writer.writerows(report_rows)


def export_excel(report_rows):
    """Export the updates to a formatted Excel workbook."""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Daily Updates"

    worksheet.append(HEADERS)

    for row in report_rows:
        worksheet.append(row)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    column_widths = {
        "A": 14,
        "B": 14,
        "C": 20,
        "D": 20,
        "E": 45,
        "F": 28,
        "G": 28,
        "H": 32,
        "I": 25,
        "J": 65,
        "K": 20,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    workbook.save(EXCEL_REPORT_PATH)


if __name__ == "__main__":
    database_rows = get_update_rows()
    report_rows = build_report_rows(database_rows)

    export_csv(report_rows)
    export_excel(report_rows)

    print(f"Exported {len(report_rows)} update(s).")
    print(f"CSV report: {CSV_REPORT_PATH}")
    print(f"Excel report: {EXCEL_REPORT_PATH}")